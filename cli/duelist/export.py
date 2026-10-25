import os
import sys
from datetime import date, timedelta
from typing import List, Dict, Any, Tuple, Union


def categorize_assignments(
    assignments: Union[List[Dict[str, Any]], Dict[str, Any]],
    today: date = None
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Categorize assignments into four distinct sections:
    1. Overdue (incomplete, due < today)
    2. Due This Week (incomplete, today <= due <= today + 7 days)
    3. Coming Up Later (incomplete, due > today + 7 days)
    4. Completed (done is True)
    """
    if isinstance(assignments, dict) and "assignments" in assignments:
        items = assignments["assignments"]
    elif isinstance(assignments, list):
        items = assignments
    else:
        items = []

    if today is None:
        today = date.today()

    cutoff = today + timedelta(days=7)

    overdue = []
    due_this_week = []
    coming_up_later = []
    completed = []

    for a in sorted(items, key=lambda x: str(x.get("due", ""))):
        if a.get("done", False):
            completed.append(a)
            continue

        due_val = a.get("due")
        if not due_val:
            coming_up_later.append(a)
            continue

        try:
            due_date = date.fromisoformat(due_val)
        except (ValueError, TypeError):
            coming_up_later.append(a)
            continue

        if due_date < today:
            overdue.append(a)
        elif today <= due_date <= cutoff:
            due_this_week.append(a)
        else:
            coming_up_later.append(a)

    return overdue, due_this_week, coming_up_later, completed


def generate_summary(assignments: Union[List[Dict[str, Any]], Dict[str, Any]], output_path: str = None) -> str:
    """Generate weekly summary docx file using python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        sys.stderr.write("Missing dependency python-docx. Run: pip install python-docx\n")
        sys.exit(1)

    today = date.today()
    overdue, this_week, later, completed = categorize_assignments(assignments, today)

    doc = Document()

    title = doc.add_heading("DueList Weekly Summary", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(f"Generated on {today.strftime('%B %d, %Y')}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle.runs:
        subtitle.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        subtitle.runs[0].font.size = Pt(11)

    total_count = len(overdue) + len(this_week) + len(later) + len(completed)
    summary_p = doc.add_paragraph(
        f"Summary: {total_count} Total · {len(overdue)} Overdue · {len(this_week)} Due This Week · "
        f"{len(later)} Coming Up Later · {len(completed)} Completed"
    )
    summary_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if summary_p.runs:
        summary_p.runs[0].font.italic = True
        summary_p.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    def add_section(heading: str, task_list: List[Dict[str, Any]], color: RGBColor):
        h = doc.add_heading(heading, level=1)
        if not task_list:
            p = doc.add_paragraph("Nothing here.")
            if p.runs:
                p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                p.runs[0].italic = True
        else:
            for a in task_list:
                p = doc.add_paragraph(style="List Bullet")
                title_run = p.add_run(a.get("title", "Untitled"))
                title_run.bold = True
                title_run.font.color.rgb = color

                cls_str = a.get("class", "").strip()
                cls_display = f" — {cls_str}" if cls_str else ""

                due_str = a.get("due", "")
                try:
                    formatted_due = date.fromisoformat(due_str).strftime("%B %d, %Y")
                except Exception:
                    formatted_due = due_str

                p.add_run(f"{cls_display} | Due {formatted_due}")

    add_section("Overdue", overdue, RGBColor(0xFF, 0x6B, 0x6B))
    add_section("Due This Week", this_week, RGBColor(0xE6, 0x9A, 0x00))
    add_section("Coming Up Later", later, RGBColor(0x4A, 0x4E, 0x69))
    add_section("Completed", completed, RGBColor(0x2E, 0x9E, 0x4F))

    if output_path is None:
        output_path = os.path.abspath(f"DueList_Summary_{today.isoformat()}.docx")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path


def generate_summary_md(assignments: Union[List[Dict[str, Any]], Dict[str, Any]], output_path: str = None) -> str:
    """Generate weekly summary Markdown file."""
    today = date.today()
    overdue, this_week, later, completed = categorize_assignments(assignments, today)

    def format_section(heading: str, task_list: List[Dict[str, Any]]) -> List[str]:
        lines = [f"## {heading}", ""]
        if not task_list:
            lines.append("_Nothing here._")
            lines.append("")
        else:
            for a in task_list:
                title = a.get("title", "Untitled")
                cls_str = a.get("class", "").strip()
                cls_display = f" — {cls_str}" if cls_str else ""
                due_str = a.get("due", "")
                try:
                    formatted_due = date.fromisoformat(due_str).strftime("%B %d, %Y")
                except Exception:
                    formatted_due = due_str
                lines.append(f"- **{title}**{cls_display} | Due {formatted_due}")
            lines.append("")
        return lines

    total_count = len(overdue) + len(this_week) + len(later) + len(completed)

    lines = [
        "# DueList Weekly Summary",
        f"_Generated on {today.strftime('%B %d, %Y')}_",
        "",
        f"**Overview:** {total_count} Total · {len(overdue)} Overdue · {len(this_week)} Due This Week · "
        f"{len(later)} Coming Up Later · {len(completed)} Completed",
        "",
    ]
    lines.extend(format_section("Overdue", overdue))
    lines.extend(format_section("Due This Week", this_week))
    lines.extend(format_section("Coming Up Later", later))
    lines.extend(format_section("Completed", completed))

    content = "\n".join(lines)

    if output_path is None:
        output_path = os.path.abspath(f"DueList_Summary_{today.isoformat()}.md")
    else:
        output_path = os.path.abspath(output_path)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)

    return output_path
